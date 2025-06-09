import pandas as pd
import io
import zipfile
from openpyxl.utils import get_column_letter
from flask import send_file
from app import db
from app.models.inventory import Item, InventoryRecord, SalesPlatform, Listing, SaleTransaction

def detect_file_type(file):
    """Detect if file is Excel or CSV based on filename"""
    filename = file.filename.lower()
    if filename.endswith('.xlsx') or filename.endswith('.xls'):
        return 'excel'
    elif filename.endswith('.csv'):
        return 'csv'
    else:
        return None

def read_file(file):
    """Read Excel or CSV file into pandas DataFrame"""
    file_type = detect_file_type(file)
    if file_type == 'excel':
        return pd.read_excel(file)
    elif file_type == 'csv':
        return pd.read_csv(file)
    else:
        raise ValueError("Unsupported file type. Please upload an Excel or CSV file.")

def export_data(data_type, file_format='excel'):
    """Export data to Excel or CSV"""
    # Get data based on type
    if data_type == 'inventory':
        df = get_inventory_data()
    elif data_type == 'platforms':
        df = get_platforms_data()
    elif data_type == 'listings':
        df = get_listings_data()
    elif data_type == 'transactions':
        df = get_transactions_data()
    elif data_type == 'all':
        return export_all_data(file_format)
    else:
        raise ValueError(f"Unknown data type: {data_type}")
    
    # Export to requested format
    output = io.BytesIO()
    if file_format == 'excel':
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name=data_type.capitalize(), index=False)
            # Auto-adjust columns' width
            for column in df:
                column_width = max(df[column].astype(str).map(len).max(), len(column))
                col_idx = df.columns.get_loc(column)
                writer.sheets[data_type.capitalize()].column_dimensions[get_column_letter(col_idx + 1)].width = column_width + 2
        mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        filename = f"{data_type}_export.xlsx"
    else:  # CSV
        df.to_csv(output, index=False)
        mimetype = 'text/csv'
        filename = f"{data_type}_export.csv"
    
    output.seek(0)
    return send_file(output, mimetype=mimetype, as_attachment=True, download_name=filename)

def export_all_data(file_format='excel'):
    """Export all data to a single file with multiple sheets (Excel) or multiple files (CSV)"""
    if file_format == 'excel':
        # For Excel, we can use multiple sheets
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            # Add each data type to a separate sheet
            dfs = {
                'Items': get_inventory_data(),
                'Platforms': get_platforms_data(),
                'Listings': get_listings_data(),
                'Transactions': get_transactions_data()
            }
            
            for sheet_name, df in dfs.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                # Auto-adjust columns' width
                for column in df:
                    column_width = max(df[column].astype(str).map(len).max(), len(column))
                    col_idx = df.columns.get_loc(column)
                    writer.sheets[sheet_name].column_dimensions[get_column_letter(col_idx + 1)].width = column_width + 2
        
        output.seek(0)
        return send_file(
            output, 
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name="inventory_system_export.xlsx"
        )
    else:
        # For CSV, we need to create a zip file with multiple CSVs
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as zipf:
            # Add each data type to a separate CSV file in the zip
            dfs = {
                'items': get_inventory_data(),
                'platforms': get_platforms_data(),
                'listings': get_listings_data(),
                'transactions': get_transactions_data()
            }
            
            for name, df in dfs.items():
                csv_data = io.StringIO()
                df.to_csv(csv_data, index=False)
                zipf.writestr(f"{name}_export.csv", csv_data.getvalue())
        
        output.seek(0)
        return send_file(
            output,
            mimetype='application/zip',
            as_attachment=True,
            download_name="inventory_system_export.zip"
        )

# Helper functions to get data for each model
def get_inventory_data():
    """Get inventory items and records as a DataFrame"""
    items = Item.query.all()
    data = []
    for item in items:
        # Add a row for each inventory record
        records = item.inventory_records.all()
        if records:
            for record in records:
                data.append({
                    'Item ID': item.id,
                    'Name': item.name,
                    'SKU': item.sku,
                    'Description': item.description,
                    'Condition': item.condition,
                    'Cost Price': item.cost_price,
                    'Record ID': record.id,
                    'Quantity': record.quantity,
                    'Location': record.location,
                    'Date Added': record.date_added
                })
        else:
            # Add a row for the item even if it has no records
            data.append({
                'Item ID': item.id,
                'Name': item.name,
                'SKU': item.sku,
                'Description': item.description,
                'Condition': item.condition,
                'Cost Price': item.cost_price,
                'Record ID': None,
                'Quantity': 0,
                'Location': None,
                'Date Added': None
            })
    
    return pd.DataFrame(data)

def get_platforms_data():
    """Get sales platforms as a DataFrame"""
    platforms = SalesPlatform.query.all()
    data = [{
        'Platform ID': platform.id,
        'Name': platform.name,
        'URL': platform.url,
        'Fee Percentage': platform.fee_percentage
    } for platform in platforms]
    
    return pd.DataFrame(data)

def get_listings_data():
    """Get listings as a DataFrame"""
    listings = Listing.query.all()
    data = [{
        'Listing ID': listing.id,
        'Item ID': listing.item_id,
        'Item Name': listing.item.name,
        'Item SKU': listing.item.sku,
        'Platform ID': listing.platform_id,
        'Platform Name': listing.platform.name,
        'Listing URL': listing.listing_url,
        'Listing Price': listing.listing_price,
        'Status': listing.status,
        'Date Listed': listing.date_listed,
        'Date Sold': listing.date_sold
    } for listing in listings]
    
    return pd.DataFrame(data)

def get_transactions_data():
    """Get sales transactions as a DataFrame"""
    transactions = SaleTransaction.query.all()
    data = [{
        'Transaction ID': transaction.id,
        'Listing ID': transaction.listing_id,
        'Item Name': transaction.listing.item.name,
        'Item SKU': transaction.listing.item.sku,
        'Platform': transaction.listing.platform.name,
        'Final Sale Price': transaction.final_sale_price,
        'Tax Amount': transaction.tax_amount,
        'Platform Fee': transaction.platform_fee,
        'Shipping Cost': transaction.shipping_cost,
        'Other Fees': transaction.other_fees,
        'Date Sold': transaction.date_sold,
        'Profit': transaction.calculate_profit(),
        'Profit Margin': transaction.calculate_profit_margin()
    } for transaction in transactions]
    
    return pd.DataFrame(data)

def import_inventory(file):
    """Import inventory items and records from Excel or CSV"""
    try:
        df = read_file(file)
        
        # Minimal validation - just check if we have either Name or SKU
        if 'Name' not in df.columns and 'SKU' not in df.columns:
            return False, "File must contain at least 'Name' or 'SKU' column"
        
        results = {
            'success': 0,
            'updated': 0,
            'errors': [],
            'total': len(df)
        }
        
        # Process each row
        for index, row in df.iterrows():
            try:
                # Check if we can identify the item
                item = None
                if 'SKU' in df.columns and pd.notna(row.get('SKU')):
                    item = Item.query.filter_by(sku=row['SKU']).first()
                
                if item:
                    # Update existing item with any provided fields
                    if 'Name' in df.columns and pd.notna(row.get('Name')):
                        item.name = row['Name']
                    if 'Description' in df.columns and pd.notna(row.get('Description')):
                        item.description = row['Description']
                    if 'Condition' in df.columns and pd.notna(row.get('Condition')):
                        item.condition = row['Condition']
                    if 'Cost Price' in df.columns and pd.notna(row.get('Cost Price')):
                        item.cost_price = float(row['Cost Price'])
                    
                    results['updated'] += 1
                else:
                    # Create new item
                    if 'Name' not in df.columns or pd.isna(row.get('Name')):
                        results['errors'].append(f"Row {index+2}: Missing required field 'Name' for new item")
                        continue
                    
                    item = Item(
                        name=row['Name'],
                        sku=row.get('SKU') if 'SKU' in df.columns and pd.notna(row.get('SKU')) else None,
                        description=row.get('Description') if 'Description' in df.columns and pd.notna(row.get('Description')) else None,
                        condition=row.get('Condition') if 'Condition' in df.columns and pd.notna(row.get('Condition')) else None,
                        cost_price=float(row['Cost Price']) if 'Cost Price' in df.columns and pd.notna(row.get('Cost Price')) else 0.0
                    )
                    db.session.add(item)
                    db.session.flush()  # Get the item ID
                    results['success'] += 1
                
                # Add inventory record if quantity is provided
                if 'Quantity' in df.columns and pd.notna(row.get('Quantity')) and float(row['Quantity']) > 0:
                    record = InventoryRecord(
                        item_id=item.id,
                        quantity=int(float(row['Quantity'])),
                        location=row.get('Location') if 'Location' in df.columns and pd.notna(row.get('Location')) else None
                    )
                    db.session.add(record)
                
            except Exception as e:
                results['errors'].append(f"Row {index+2}: {str(e)}")
        
        if len(results['errors']) < results['total']:  # If at least some rows were processed successfully
            db.session.commit()
        else:
            db.session.rollback()
        
        return True, results
    
    except Exception as e:
        db.session.rollback()
        return False, str(e)

def import_platforms(file):
    """Import sales platforms from Excel or CSV"""
    try:
        df = read_file(file)
        
        # Minimal validation
        if 'Name' not in df.columns:
            return False, "File must contain 'Name' column"
        
        results = {
            'success': 0,
            'updated': 0,
            'errors': [],
            'total': len(df)
        }
        
        # Process each row
        for index, row in df.iterrows():
            try:
                # Check if platform exists
                platform = SalesPlatform.query.filter_by(name=row['Name']).first()
                
                if platform:
                    # Update existing platform
                    if 'URL' in df.columns and pd.notna(row.get('URL')):
                        platform.url = row['URL']
                    if 'Fee Percentage' in df.columns and pd.notna(row.get('Fee Percentage')):
                        platform.fee_percentage = float(row['Fee Percentage'])
                    
                    results['updated'] += 1
                else:
                    # Create new platform
                    platform = SalesPlatform(
                        name=row['Name'],
                        url=row.get('URL') if 'URL' in df.columns and pd.notna(row.get('URL')) else None,
                        fee_percentage=float(row['Fee Percentage']) if 'Fee Percentage' in df.columns and pd.notna(row.get('Fee Percentage')) else 0.0
                    )
                    db.session.add(platform)
                    results['success'] += 1
                
            except Exception as e:
                results['errors'].append(f"Row {index+2}: {str(e)}")
        
        if len(results['errors']) < results['total']:  # If at least some rows were processed successfully
            db.session.commit()
        else:
            db.session.rollback()
        
        return True, results
    
    except Exception as e:
        db.session.rollback()
        return False, str(e)

def create_sample_template(data_type, file_format='excel'):
    """Create a sample template file for import"""
    if data_type == 'inventory':
        # Create sample inventory template
        df = pd.DataFrame({
            'Name': ['Sample Item 1', 'Sample Item 2'],
            'SKU': ['SKU001', 'SKU002'],
            'Description': ['This is a sample item', 'Another sample item'],
            'Condition': ['new', 'used'],
            'Cost Price': [10.99, 5.99],
            'Quantity': [5, 10],
            'Location': ['Warehouse A', 'Warehouse B']
        })
    elif data_type == 'platforms':
        # Create sample platforms template
        df = pd.DataFrame({
            'Name': ['eBay', 'Amazon', 'Etsy'],
            'URL': ['https://www.ebay.com', 'https://www.amazon.com', 'https://www.etsy.com'],
            'Fee Percentage': [10.0, 15.0, 5.0]
        })
    else:
        raise ValueError(f"Unknown template type: {data_type}")
    
    # Export to requested format
    output = io.BytesIO()
    if file_format == 'excel':
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Template', index=False)
            # Auto-adjust columns' width
            for column in df:
                column_width = max(df[column].astype(str).map(len).max(), len(column))
                col_idx = df.columns.get_loc(column)
                writer.sheets['Template'].column_dimensions[get_column_letter(col_idx + 1)].width = column_width + 2
        mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        filename = f"{data_type}_template.xlsx"
    else:  # CSV
        df.to_csv(output, index=False)
        mimetype = 'text/csv'
        filename = f"{data_type}_template.csv"
    
    output.seek(0)
    return send_file(output, mimetype=mimetype, as_attachment=True, download_name=filename)
